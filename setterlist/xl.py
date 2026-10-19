import csv, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

rows = list(csv.reader(open("juan-call-list.csv")))
head, body = rows[0], rows[1:]
wb = Workbook()
ws = wb.active; ws.title = "Call list"
ws.append(head)
NUMCOL = {0, 6, 7, 8}
for r in body:
    ws.append([int(v) if (i in NUMCOL and str(v).strip().lstrip("-").isdigit()) else v
               for i, v in enumerate(r)])

hfill = PatternFill("solid", fgColor="1F3864")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = hfill
    c.alignment = Alignment(vertical="center", wrap_text=True)
ws.freeze_panes = "A2"; ws.row_dimensions[1].height = 30
bands = {"A - call first": "C6E0B4", "B - strong": "FFE699", "C - worth a dial": "F2F2F2"}
for i, r in enumerate(body, start=2):
    fill = bands.get(r[1])
    if fill:
        ws.cell(row=i, column=2).fill = PatternFill("solid", fgColor=fill)
    ws.cell(row=i, column=6).alignment = Alignment(wrap_text=True, vertical="top")
widths = [6, 16, 24, 16, 30, 62, 7, 10, 11, 20, 26, 30, 18, 24, 12, 12, 22, 40]
for n, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(n)].width = w
ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{len(body)+1}"

m = wb.create_sheet("How this list was built")
tag = json.load(open("tag_result.json"))
lines = [
 ("Setter call list — built 19 October 2026", True),
 ("", False),
 (f"250 people, ranked. All but 2 live in the SMS sub-account and carry the tag \"{tag['tag']}\" ({tag['tagged']} contacts), so Juan can pull the list up in GHL as a smart list or filter.", False),
 ("", False),
 ("Where it came from", True),
 ("Every contact in both sub-accounts was read: 48,695 in SMS and 49,092 in the main/EMAIL account, plus their conversations, opportunities and calendar appointments. People who appear in both accounts were matched on email and phone and kept once, with the SMS record as the one Juan calls.", False),
 ("", False),
 ("Who was taken out", True),
 ("Already owns a product — 2,491 people. Ownership was read from purchase and member tags (options navigator, options bot, futures, VTF access, uoo member, customer, won), from won opportunities, from the Bot / Futures / Flight Path support pipelines, and from the survey question asking which programs they own.", False),
 ("No number anyone can dial — 12,767 people. Blank, non-US, or malformed numbers.", False),
 ("Do-not-call — 11,109 people. Global DND, call-channel DND, or a dnd / bad number / verified bad / spamtrap tag. An SMS-only or email-only opt-out does not exclude anyone, since the phone is still fair game.", False),
 ("Already has an appointment coming up — 54 people. Confirmed future bookings on every calendar in both accounts, plus anyone whose opportunity is sitting in a booked or scheduled stage changed in the last 45 days. Each of the final 250 was then re-checked one by one against GHL's own appointment record: none of them has anything on the books.", False),
 ("Disqualified — 6,226 people. The Disqualified / Bad Lead / Unqualified stages and tags. People who simply never finished a survey were left in.", False),
 ("No sign of life — 38,639 people. Nothing new, nothing engaged: no recent lead date, no reply, no webinar, no click, no activity.", False),
 ("", False),
 ("How \"most likely to book\" was scored", True),
 ("Points add up across five signals, so the people at the top are both warm and current:", False),
 ("Freshness of the lead — up to 45 points. A lead 3 days old scores 45, a week 38, a month 22. New leads answer the phone.", False),
 ("They have actually written back — up to 40 points. A reply in the last week scores 40, within the month 32, within three months 20. This is the single strongest signal of someone who will pick up.", False),
 ("Webinar behaviour — up to 34 points. Attending beats watching the replay, which beats registering and not showing. The points decay after 45 days and again after 120.", False),
 ("Clicks and opens — up to 12 points, plus 14 for a hot / warm / interested / replied tag.", False),
 ("Live in a pipeline — up to 14 points for an open opportunity sitting in New Lead, Responded, Engaged or Follow Up.", False),
 ("A prior no-show or cancel costs 6 points; it is a reason to call, not a reason to skip.", False),
 ("", False),
 ("Priority A is 70 points or more (115 people), B is 55-69 (95 people), C is below 55 (40 people).", False),
 ("Sixty of the 250 places were held for the newest leads of the last 30 days, which is why some of them sit lower on score: 30 of the list came in within the past week, 84 within the past month, and 146 have replied to us within the last 30 days.", False),
 ("", False),
 ("Worth knowing", True),
 ("Two of the 250 exist only in the main/EMAIL account, so they could not be tagged in SMS. They are rows marked Main/Email in the sub-account column.", False),
 ("Nothing else in GHL was changed. The only write was the tag.", False),
]
for text, bold in lines:
    m.append([text]); m.cell(row=m.max_row, column=1).font = Font(bold=bold, size=12 if bold else 11)
    m.cell(row=m.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
m.column_dimensions["A"].width = 120
wb.save("/mnt/project-files/setter-lists/juan-call-list-2026-10-19.xlsx")
print("saved xlsx")
